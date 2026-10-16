import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Sets cell padding in dxa."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def generate_document():
    doc = Document()

    # Set Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Base typography styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # -------------------------------------------------------------------------
    # COVER TITLE & METADATA
    # -------------------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(14)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("GEO AUCTION & REAL-TIME BIDDING SYSTEM")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(4, 120, 87) # Emerald 700

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Comprehensive Architectural Specification, Database Models & Line-by-Line Code Walkthrough")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_info = [
        ("Platform Version:", "Geo Auction System v3.2 Enterprise (Bayut Luxury Standard)"),
        ("Architecture Model:", "Full-Stack PHP 8.x (OOP & PDO) + MySQL + Vanilla JS (AJAX) + Tailwind CSS"),
        ("Primary Live File:", "live_auction.php (Complete Real-Time Bidding Arena & Showcase)"),
        ("Document Purpose:", "Exhaustive Working Explanation, Line-by-Line Code Breakdown & Production Manual")
    ]

    for r_idx, row in enumerate(meta_table.rows):
        label, val = meta_info[r_idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, 80, 80, 140, 140)
        set_cell_margins(cell_val, 80, 80, 140, 140)

        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(15, 23, 42)

        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------------------
    # HELPER FUNCTIONS
    # -------------------------------------------------------------------------
    def add_heading(text, level=1):
        p = doc.add_paragraph()
        if level == 1:
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(4, 120, 87)
        elif level == 2:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.bold = True
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(15, 23, 42)
        else:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 41, 59)

    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix + " ")
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(15, 23, 42)
        r_t = p.add_run(text)
        r_t.font.italic = italic
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        rb = p.add_run(bold_prefix + ": ")
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(15, 23, 42)
        rt = p.add_run(text)
        return p

    def add_code_block(code_str):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "0F172A") # Slate 900
        set_cell_margins(cell, 100, 100, 140, 140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(code_str)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(241, 245, 249) # Slate 100
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_callout(text, title="CRITICAL ARCHITECTURE NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "F0FDF4") # Emerald 50
        set_cell_margins(cell, 100, 100, 140, 140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"[{title}]\n")
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(4, 120, 87)
        r_msg = p.add_run(text)
        r_msg.font.size = Pt(9.5)
        r_msg.font.color.rgb = RGBColor(22, 101, 52)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------------------
    # CHAPTER 1: EXECUTIVE OVERVIEW & PLATFORM ARCHITECTURE
    # -------------------------------------------------------------------------
    add_heading("1. Executive Overview & System Architecture", level=1)
    add_p("The Geo Auction & Real-Time Bidding System is a high-performance web platform combining location-based item discovery, multi-agent bidding algorithms, luxury real estate presentation standards (modeled after Bayut Dubai), and bank-grade escrow wallet settlement. Unlike traditional e-commerce platforms that rely on static database reads, this application functions as a reactive auction floor where countdown timers, bid streams, escrow holds, and anti-sniping extensions operate in real-time without requiring full page reloads.")

    add_heading("1.1 Core Technology Stack & Design Principles", level=2)
    add_bullet("Backend Runtime", "PHP 8.2+ utilizing Strict Object-Oriented Programming (OOP) with prepared PDO statements, strict parameter binding, and granular transaction rollbacks.")
    add_bullet("Database Architecture", "MySQL 8.0 / MariaDB with InnoDB engine, enforcing foreign key integrity, row-level locking (FOR UPDATE) for concurrency, and optimized indexes.")
    add_bullet("Frontend Presentation Layer", "Tailwind CSS v3 (scoped utility classes) combined with custom Vanilla CSS for hardware-accelerated animations (Ken Burns motion, cubic-bezier transforms, and micro-pulse indicators).")
    add_bullet("Asynchronous Client Layer", "Pure Vanilla JavaScript with standard Fetch API & XMLHttpRequest (AJAX), ensuring 60 FPS UI rendering, zero external bloat, and full cross-browser compatibility.")
    add_bullet("Geospatial & Mapping Engine", "Leaflet.js open-source map integration with Haversine spherical distance calculation (in kilometers) to sort lots closest to the user's GPS/IP location.")

    add_heading("1.2 Repository Directory Structure & File Map", level=2)
    add_p("The project is structured according to clear separation of concerns (Presentation, Business Logic, Database Access, and Asynchronous Endpoints):")
    add_code_block("""geo-auction-system-GIT/
├── includes/
│   ├── Database.php          -> Singleton PDO & MySQLi connection manager
│   ├── AuctionEngine.php      -> Bidding mechanics (Popcorn, Dutch, Vickrey)
│   ├── EscrowService.php      -> Pre-bid wallet hold and fund release engine
│   ├── SecondChanceService.php-> Automated 24h runner-up offer generator
│   ├── auth.php              -> Authentication, CSRF, and session security
│   ├── header.php            -> Universal global navigation bar and balance
│   └── footer.php            -> System footer, copyright, and global scripts
├── api/
│   ├── agent.php             -> TruBroker REST endpoint & AJAX click tracking
│   └── wallet.php            -> Wallet balance and transaction history API
├── ajax/
│   ├── comments.php          -> Real-time lot discussions and question feed
│   ├── dutch.php             -> Dutch auction price decay ticker
│   └── reaction.php          -> Emoji reaction stream for live auction items
├── uploads/
│   ├── items/                -> Photorealistic listing photos (e.g. mudon_*.jpg)
│   ├── broker_asif.png       -> Authentic TruBroker profile portrait
│   └── ark_logo.png          -> Certified Real Estate Agency badge
├── live_auction.php          -> Primary Live Arena & Interactive Showcase
├── login.php & register.php  -> Hardened auth with rate limiting and CSRF
├── index.php                 -> Public marketplace homepage and hero banner
└── config/database.php       -> Centralized database credentials (geo_auction_db)""")

    # -------------------------------------------------------------------------
    # CHAPTER 2: DATABASE SCHEMA & DATA RELATIONSHIPS
    # -------------------------------------------------------------------------
    add_heading("2. Database Schema & Relational Models", level=1)
    add_p("The database 'geo_auction_db' contains 37 relational tables designed to handle high transaction concurrency while preventing race conditions. Below are the key tables powering the live auction and broker systems:")

    add_heading("2.1 Table: agents (TruBroker™ Verified Real Estate Profiles)", level=2)
    add_p("Stores certified broker profiles, verification statuses, and direct contact endpoints:")
    add_code_block("""CREATE TABLE `agents` (
    `id` INT AUTO_INCREMENT PRIMARY KEY,
    `name` VARCHAR(150) NOT NULL,
    `avatar_url` VARCHAR(255) DEFAULT NULL,
    `banner_url` VARCHAR(255) DEFAULT NULL,
    `email` VARCHAR(150) DEFAULT NULL,
    `phone` VARCHAR(50) DEFAULT NULL,
    `whatsapp_number` VARCHAR(50) DEFAULT NULL,
    `is_quality_lister` TINYINT(1) DEFAULT 1,
    `is_responsive_broker` TINYINT(1) DEFAULT 1,
    `agency_name` VARCHAR(150) DEFAULT NULL,
    `agency_logo` VARCHAR(255) DEFAULT NULL,
    `agency_properties_url` VARCHAR(255) DEFAULT NULL,
    `created_at` TIMESTAMP DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;""")

    add_heading("2.2 Table: featured_communities (Bayut Community Sub-Card)", level=2)
    add_p("Powers the inner rounded community recommendation box linked to specific brokers:")
    add_code_block("""CREATE TABLE `featured_communities` (
    `id` INT AUTO_INCREMENT PRIMARY KEY,
    `agent_id` INT NOT NULL,
    `title` VARCHAR(150) NOT NULL,
    `description` VARCHAR(255) DEFAULT NULL,
    `image_url` VARCHAR(255) DEFAULT NULL,
    `target_url` VARCHAR(255) DEFAULT NULL,
    `created_at` TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    INDEX (`agent_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;""")

    add_heading("2.3 Table: agent_interactions (Analytics & Engagement Tracking)", level=2)
    add_p("Records asynchronous user clicks on Email, Call, WhatsApp, and Portfolio triggers for lead analytics:")
    add_code_block("""CREATE TABLE `agent_interactions` (
    `id` INT AUTO_INCREMENT PRIMARY KEY,
    `agent_id` INT NOT NULL,
    `interaction_type` ENUM('email', 'call', 'whatsapp', 'view_properties') NOT NULL,
    `user_ip` VARCHAR(45) DEFAULT NULL,
    `created_at` TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    INDEX (`agent_id`),
    INDEX (`interaction_type`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;""")

    add_heading("2.4 Table: items, bids, and item_images (Core Auction Relational Triad)", level=2)
    add_p("The item catalog is decoupled from images to permit unlimited high-resolution photos per listing:")
    add_bullet("items Table", "Stores auction lot title, starting_price, current_price, reserve_price, end_time, extended_time (anti-sniping timestamp), category_id, seller_id, and GPS latitude/longitude.")
    add_bullet("item_images Table", "Stores item_id, image_path, sort_order, and is_featured flag. The OOP AuctionImageRepository queries this table to construct high-res showcase galleries.")
    add_bullet("bids Table", "Stores item_id, user_id, bidder_name, bid_amount, is_winning, bid_status, and created_at with millisecond precision.")

    # -------------------------------------------------------------------------
    # CHAPTER 3: ADVANCED AUCTION MECHANICS & CORE ALGORITHMS
    # -------------------------------------------------------------------------
    add_heading("3. Advanced Bidding Mechanics & Algorithms", level=1)
    add_p("The platform introduces enterprise bidding mechanisms designed to maximize fair price discovery, stop automated sniper bots, and secure seller settlements:")

    add_heading("3.1 Anti-Sniping (Soft Close / Popcorn Bidding)", level=2)
    add_p("In traditional online auctions, malicious bots submit bids in the final 2 seconds ('bid sniping') preventing legitimate human buyers from responding. The Anti-Sniping engine dynamically monitors the remaining duration:")
    add_callout(
        "If a bid is received when remaining time <= 120 seconds (2 minutes), the auction's extended_time is immediately updated to (NOW() + 180 seconds). The frontend polling loop detects this change in real-time, displays an amber '+3 Mins Anti-Sniping Extension' alert badge, and extends the visual countdown without refreshing the page.",
        title="ANTI-SNIPING EXTENSION RULE"
    )

    add_heading("3.2 Pre-Bid Escrow Hold (EscrowService.php)", level=2)
    add_p("To prevent non-paying bidders, the bidding engine enforces pre-bid wallet locking:")
    add_bullet("Step 1 (Pre-Validation)", "When placing a bid of Rs. X, the bidder's wallet balance is verified. If balance < X, the bid is rejected instantly.")
    add_bullet("Step 2 (Fund Hold)", "The amount is transferred to 'locked_balance' via row-level locking (SELECT ... FOR UPDATE).")
    add_bullet("Step 3 (Outbid Release)", "When a higher bid of Rs. Y is submitted by another buyer, the previous bidder's locked funds are automatically released back to their active wallet balance.")
    add_bullet("Step 4 (Settlement)", "Upon auction conclusion, the winning bidder's held funds are settled directly to the seller (minus marketplace commission fees).")

    add_heading("3.3 Multi-Agent Auto-Bid (BidBot) Escalation Engine", level=2)
    add_p("Bidders can configure a maximum willingness to pay via the Auto-Bid modal. When competing bids are placed, the internal `runAutoBids()` function executes proxy escalation in increments of Rs. 100 or `min_increment` until the maximum threshold is met, releasing notifications to bidders who are outbid.")

    # -------------------------------------------------------------------------
    # CHAPTER 4: IN-DEPTH LINE-BY-LINE WALKTHROUGH OF LIVE_AUCTION.PHP
    # -------------------------------------------------------------------------
    add_heading("4. Comprehensive Line-by-Line Code Breakdown: live_auction.php", level=1)
    add_p("The primary auction arena file (live_auction.php) spans over 4,000 lines of highly optimized PHP, HTML5, CSS3, and JavaScript. Below is a detailed, section-by-section breakdown of every critical block:")

    add_heading("4.1 Lines 1–150: Session Security, CSRF, and Geo-Coordinates Initialization", level=2)
    add_bullet("Session Start & Auth Guards", "Lines 1–35 verify session validity via auth.php. Generates an encrypted CSRF token stored in session and emitted as an HTML meta tag `<meta name=\"csrf-token\">` for all outgoing AJAX requests.")
    add_bullet("Role Resolution", "Detects user roles ('guest', 'buyer', 'seller', 'admin') to conditionally permit auction management, bid acceptance, or soft-deletion controls.")
    add_bullet("User Geo-Location", "Retrieves `$_SESSION['user_lat']` and `$_SESSION['user_lng']` (defaulting to Hyderabad coordinates 25.3960, 68.3578) to compute real-time proximity in kilometers for nearby auctions.")

    add_heading("4.2 Lines 150–295: AuctionImageRepository & Multi-Source Media Fallback", level=2)
    add_p("Lines 160–225 define the class `AuctionImageRepository`. It resolves lot images from multiple storage paradigms:")
    add_code_block("""// Relational lookup -> JSON array -> Comma CSV -> Fallback Placeholder
public function getImagesByItemId(int $itemId, string $rawCsv = '', string $featured = ''): array {
    // 1. First priority: Check item_images table
    $stmt = $this->pdo->prepare("SELECT image_path FROM item_images WHERE item_id = ? ORDER BY sort_order ASC");
    $stmt->execute([$itemId]);
    $dbImages = $stmt->fetchAll(PDO::FETCH_COLUMN);
    if (!empty($dbImages)) return array_map('cleanImagePath', $dbImages);
    
    // 2. Second priority: Featured image field
    if (!empty($featured)) $images[] = cleanImagePath($featured);
    
    // 3. Third priority: CSV / JSON comma separated images
    ...
    return !empty($images) ? $images : [$this->placeholder];
}""")

    add_heading("4.3 Lines 295–720: Real-Time AJAX Dispatcher Engine", level=2)
    add_p("When `$_GET['ajax'] == '1'` is received, the script terminates standard HTML execution and acts as an ultra-fast JSON REST API:")
    add_bullet("action=poll_auction (Lines 310–440)", "Executes high-frequency polling. Queries the current highest bid, bidder name, bid count, watcher count, time remaining, and anti-sniping extension status. Returns structured JSON in < 35ms.")
    add_bullet("action=place_bid (Lines 440–525)", "Wraps bid insertion inside a database transaction (`$pdo->beginTransaction()`). Validates bid >= next_min_bid, executes escrow wallet locks, triggers anti-sniping soft-close (+3 mins) if within 120 seconds of ending, and fires `runAutoBids()`.")
    add_bullet("action=toggle_watchlist (Lines 525–560)", "Toggles items in/out of the user's `wishlist` table and returns updated watcher counts.")
    add_bullet("action=set_autobid (Lines 560–610)", "Stores user's proxy bidding ceiling in the `auto_bids` table with active flag.")
    add_bullet("action=buy_it_now (Lines 610–660)", "Executes instant purchase, sets item status to 'sold', locks the winning price, and generates buyer invoice records.")

    add_heading("4.4 Lines 935–1300: Scoped CSS Design System & Keyframe Animations", level=2)
    add_p("This section contains the pure CSS design system, responsive breakpoints, and custom animations:")
    add_bullet("@keyframes kenBurnsTour (Lines 1178–1195)", "Produces cinematic drone camera motion on listing photos by subtly translating and scaling the image from `scale(1)` to `scale(1.08)` over 12 seconds with alternate direction.")
    add_bullet("@keyframes checkedPulseRing (Lines 1150–1175)", "Produces a continuous radial breathing pulse around the green Bayut 'Checked' badge, signaling a verified, physically inspected listing.")
    add_bullet(".fs-image-modal (Lines 968–990)", "Strictly configured with `flex-direction: column !important; justify-content: space-between !important;` to ensure the modal header, zoom controls, center image viewport, and thumbnail filmstrip are aligned properly without horizontal squashing.")

    add_heading("4.5 Lines 1700–1835: Bayut Dubai Luxury Real Estate Showcase Grid", level=2)
    add_p("Renders the flagship 4-panel photography grid inspired by premier UAE luxury portals:")
    add_bullet("Main Left Column (~68% width)", "Displays `#mainImage` with instant fullscreen viewer trigger, watermarked crest overlay, pulsating Checked badge, and bottom action pills.")
    add_bullet("Bottom Action Pills (Lines 1756–1766)", "Features the `#motionTourToggleBtn` with a pulsating green status dot (`.motion-pulse-dot`), the 'Virtual Tour' trigger, and the 'View all photos' fullscreen expander.")
    add_bullet("Right Column (~32% width)", "Renders 3 stacked perspective preview cards with hover zoom and camera count badge (`<i class=\"fa-solid fa-camera\"></i> 4 Photos`).")

    add_heading("4.6 Lines 1835–1945: TruBroker™ Real Estate Agent Profile Card", level=2)
    add_p("Integrates the verified real estate broker component backed by dynamic MySQL data:")
    add_bullet("Emerald Wave Header Banner", "Deep teal-green gradient (`#104f47` to `#073834`) with geometric diagonal texture and crisp `TruBroker™` branding.")
    add_bullet("Overlapping Avatar & Name", "Circular portrait (`border-[3.5px] border-white shadow-md`) positioned at `-mt-9`, displaying broker name 'Asif Raza Khan'.")
    add_bullet("Dual Badges", "💎 Quality Lister (`#eef7fc`) and 📞 Responsive Broker (`#f8f0fa`).")
    add_bullet("3-Column Action Triggers", "Direct links for **Email** (`messages.php`), **Call** (`tel:`), and **WhatsApp** (`https://wa.me/`). Each anchor contains `data-agent-action` for non-blocking asynchronous interaction tracking.")
    add_bullet("Featured Community Sub-Card", "An inner rounded container showcasing community thumbnail (`mudon_aerial.jpg`), community title (`Mudon Al Ranim 8`), attractions subtitle, and right chevron.")

    add_heading("4.7 Lines 2450–2650: Interactive Lightbox Viewer & 4K Virtual Walkthrough Modal", level=2)
    add_bullet("Fullscreen Lightbox Modal (`#fullscreenImageViewerModal`)", "Complete modal overlay with real-time zoom percentage ticker (100% to 500%), smooth drag-to-pan boundary clamping, previous/next keyboard navigation, and synchronized thumbnail strip.")
    add_bullet("Virtual Walkthrough Modal (`#videoTourModal`)", "Interactive 4K inspection player with live HDR status badge, animated timeline progress bar, perspective room switchers ('Exterior Aerial', 'Living Lounge', 'Master Suite', 'Gourmet Kitchen'), and broker callback request form.")

    add_heading("4.8 Lines 2650–3200: Vanilla JavaScript Gallery & Touch Gesture Engine", level=2)
    add_p("Contains the gesture math and transformation logic for interactive photo inspection:")
    add_bullet("applyZoomTransform()", "Applies hardware-accelerated CSS `translate3d(X, Y, 0) scale(S)` to ensure smooth 60 FPS panning.")
    add_bullet("clampPanBounds()", "Strict boundary math ensuring zoomed images never get dragged outside the viewport boundary.")
    add_bullet("bindTouchGestures()", "Calculates finger distance for pinch-to-zoom and touch delta for left/right swipe image navigation on mobile devices.")

    add_heading("4.9 Lines 3200–3980: Polling Loop, Audio Chimes, and Outbid Alerts", level=2)
    add_bullet("pollAuctionData()", "Executes recursive polling every 2,500ms. If current price increases, flashes green ticker highlight and plays Web Audio API chime tone.")
    add_bullet("Anti-Sniping Countdown Engine", "Computes remaining milliseconds: `Math.floor((targetMs - now) / 1000)`. Updates hours, minutes, and seconds counters. If extended by anti-sniping, updates countdown end-time seamlessly.")

    add_heading("4.10 Lines 3980–4126: TruBroker™ Interaction Tracking AJAX & Initialization", level=2)
    add_p("Handles DOM lifecycle and asynchronous interaction reporting:")
    add_code_block("""function initAgentProfileCard() {
    var card = document.getElementById('realEstateAgentCard');
    if (!card) return;
    var agentId = card.getAttribute('data-agent-id') || 1;
    var actionBtns = card.querySelectorAll('.agent-contact-btn');
    
    actionBtns.forEach(function(btn) {
        btn.addEventListener('click', function(e) {
            var action = this.getAttribute('data-agent-action');
            if (!action) return;
            btn.classList.add('scale-95');
            setTimeout(() => btn.classList.remove('scale-95'), 180);
            
            // Asynchronous tracking call to api/agent.php
            var fd = new FormData();
            fd.append('agent_id', agentId);
            fd.append('interaction_type', action);
            fetch('api/agent.php', { method: 'POST', body: fd }).catch(function(){});
        });
    });
}""")

    # -------------------------------------------------------------------------
    # CHAPTER 5: SECURITY HARDENING & PRODUCTION AUDIT
    # -------------------------------------------------------------------------
    add_heading("5. Security Hardening & Application Auditing", level=1)
    add_bullet("SQL Injection Immunity", "All database operations employ prepared PDO statements with parameter placeholders (`?` or `:name`). Zero raw string concatenation in queries.")
    add_bullet("Cross-Site Scripting (XSS) Mitigation", "All user-supplied and database-retrieved strings are escaped using `htmlspecialchars(..., ENT_QUOTES, 'UTF-8')` prior to HTML output.")
    add_bullet("Cross-Site Request Forgery (CSRF)", "State-changing POST actions (placing bids, buy-it-now, auto-bid setup, comments) validate cryptographic session tokens.")
    add_bullet("Brute-Force & Session Protection", "Authentication routines implement IP-based rate limiting, password hashing via `PASSWORD_BCRYPT` (cost 12), and session regeneration upon login to prevent session fixation.")

    # -------------------------------------------------------------------------
    # CHAPTER 6: STEP-BY-STEP TESTING & DEPLOYMENT GUIDE
    # -------------------------------------------------------------------------
    add_heading("6. Step-by-Step Deployment & Testing Guide", level=1)
    add_bullet("Step 1: Database Migration", "Open phpMyAdmin at http://localhost/phpmyadmin, select 'geo_auction_db', and execute the migration script provided in Chapter 2.")
    add_bullet("Step 2: Start XAMPP Services", "Ensure Apache and MySQL modules are green in the XAMPP Control Panel.")
    add_bullet("Step 3: Open Live Auction Arena", "Navigate to http://localhost/geo-auction-system-GIT/live_auction.php?id=243 in your web browser.")
    add_bullet("Step 4: Test Motion Tour", "Click 'Live Motion Tour' on the bottom-left of the main villa showcase. Observe the pulsing green indicator and 4.2-second smooth cross-fading through all 4 perspective photos with Ken Burns drone pan-and-zoom motion.")
    add_bullet("Step 5: Test Fullscreen Lightbox", "Click on the main listing photo. The fullscreen lightbox opens with the photo centered and scaled. Verify wheel zoom (+/-) and drag-to-pan capabilities.")
    add_bullet("Step 6: Test TruBroker Contact Triggers", "Click Email, Call, or WhatsApp on the broker card. Open phpMyAdmin and inspect table 'agent_interactions' to confirm that user clicks are recorded in real-time.")

    # Save to project workspace root
    output_filename = "Geo_Auction_System_Complete_Architecture_and_Code_Explanation.docx"
    output_path = os.path.join(r"c:\xampp\htdocs\geo-auction-system-GIT", output_filename)
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    generate_document()
